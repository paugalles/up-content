#!/usr/bin/env python3
import os
import sys
import re
import json
import asyncio
import concurrent.futures
from pathlib import Path
from urllib.parse import urlparse
from collections import namedtuple

# Ensure we're in the right working directory or adjust sys.path
sys.path.insert(0, str(Path(__file__).parent))

from common.articles import discover_urls, Http, fetch_article
from common.generators import EdgeTTSProvider, VideoComposer
from PIL import Image, ImageDraw, ImageFont

from google import genai
from google.genai import types

import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Setup environment variables
brand_navy = os.getenv("BRAND_COLOR_NAVY", "#0c3d6d")
brand_blue = os.getenv("BRAND_COLOR_BLUE", "#0c88eb")
brand_text = os.getenv("BRAND_COLOR_TEXT", "#0f172a")
brand_canvas = os.getenv("BRAND_COLOR_CANVAS", "#f8fafc")
brand_border = os.getenv("BRAND_COLOR_BORDER", "#e2e8f0")
cta_es = os.getenv("BRAND_CTA_ES", "Consulta la guía completa en inmibot.es.")
cta_en = os.getenv("BRAND_CTA_EN", "Read the complete guide at inmibot.es.")

def hex_to_rgb(hex_color: str):
    hex_color = hex_color.lstrip('#')
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))

Scene = namedtuple("Scene", ["spoken_text", "slide_title", "slide_body", "is_title"])

def draw_wrapped_text(draw, text, font, x, y, width_chars, fill, bullet=False, align="left"):
    import textwrap
    lines = textwrap.wrap(text, width=width_chars)
    current_y = y
    for i, line in enumerate(lines):
        bbox = draw.textbbox((0,0), line, font=font)
        line_w = bbox[2] - bbox[0]
        line_h = bbox[3] - bbox[1]
        
        draw_x = x
        if align == "center":
            draw_x = x - (line_w // 2)
            
        if bullet and i == 0:
            r = 8
            cy = current_y + (line_h // 2) + 5
            cx = draw_x - 30
            draw.ellipse([cx-r, cy-r, cx+r, cy+r], fill=fill)
            
        draw.text((draw_x, current_y), line, font=font, fill=fill)
        current_y += line_h + 15
    return current_y + 30

def clean_and_crop_image(img_path, canvas_rgb, make_transparent=False):
    img = Image.open(str(img_path)).convert("RGBA")
    pixdata = img.load()
    width, height = img.size
    corners = [pixdata[0,0], pixdata[width-1,0], pixdata[0,height-1], pixdata[width-1,height-1]]
    bg_color = max(set(corners), key=corners.count)
    
    for py in range(height):
        for px in range(width):
            r, g, b, a = pixdata[px, py]
            if abs(r - bg_color[0]) < 25 and abs(g - bg_color[1]) < 25 and abs(b - bg_color[2]) < 25:
                pixdata[px, py] = (0, 0, 0, 0)
                
    # Use the alpha channel to find the bounding box
    bbox = img.getchannel('A').getbbox()
    if bbox:
        img = img.crop(bbox)
        
    if not make_transparent:
        bg = Image.new("RGBA", img.size, canvas_rgb + (255,))
        bg.paste(img, (0,0), img)
        return bg.convert("RGB")
        
    return img

def create_pptx_and_images(slides_data, images_dir, pptx_path, slide_images_dir):
    prs = pptx.Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    blank_layout = prs.slide_layouts[6]
    
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf" if sys.platform == "darwin" else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    if not os.path.exists(font_path):
        font_path = ""

    try:
        title_font = ImageFont.truetype(font_path, 90) if font_path else ImageFont.load_default()
        subtitle_font = ImageFont.truetype(font_path, 60) if font_path else ImageFont.load_default()
        body_font = ImageFont.truetype(font_path, 40) if font_path else ImageFont.load_default()
    except OSError:
        title_font = subtitle_font = body_font = ImageFont.load_default()

    canvas_rgb = hex_to_rgb(brand_canvas)
    navy_rgb = hex_to_rgb(brand_navy)
    blue_rgb = hex_to_rgb(brand_blue)
    text_rgb = hex_to_rgb(brand_text)

    scenes = []

    # Title Slide
    title_data = slides_data["title_slide"]
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*canvas_rgb)
    
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(14), Inches(3))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_data["title"]
    p.font.size = Pt(64)
    p.font.color.rgb = RGBColor(*navy_rgb)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    img = Image.new('RGB', (1920, 1080), canvas_rgb)
    draw = ImageDraw.Draw(img)
    draw_wrapped_text(draw, title_data["title"], title_font, 960, 200, 30, navy_rgb, align="center")
    title_img_path = images_dir / "title_img.png"
    if title_img_path.exists():
        clean_title_img = clean_and_crop_image(title_img_path, canvas_rgb, make_transparent=True)
        clean_title_img_path = images_dir / "title_img_clean.png"
        clean_title_img.save(str(clean_title_img_path))
        # Place it in PPTX centered at bottom
        slide.shapes.add_picture(str(clean_title_img_path), Inches(6), Inches(4.5), width=Inches(4))
        
        # Place it in Pillow centered at bottom
        clean_title_img.thumbnail((500, 500))
        img.paste(clean_title_img, (960 - clean_title_img.width // 2, 1080 - clean_title_img.height - 100), clean_title_img)
        
    img.save(slide_images_dir / "slide_0.png")
    scenes.append(Scene(title_data["spoken_text"], title_data["title"], "", True))
    
    # Content slides
    for i, c_data in enumerate(slides_data["content_slides"]):
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*canvas_rgb)
        
        txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(14), Inches(1.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = c_data["title"]
        p.font.size = Pt(44)
        p.font.color.rgb = RGBColor(*blue_rgb)
        p.font.bold = True
        
        txBox2 = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6.5), Inches(6))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        
        bullets = c_data.get("text_left", "")
        if isinstance(bullets, str):
            bullets = bullets.split('\n')
        elif not isinstance(bullets, list):
            bullets = [str(bullets)]
            
        for bullet in bullets:
            if not str(bullet).strip(): continue
            p = tf2.add_paragraph()
            p.text = str(bullet).strip().lstrip('-').strip()
            p.font.size = Pt(28)
            p.font.color.rgb = RGBColor(*text_rgb)
            p.level = 0
            
        img_path = images_dir / f"content_img_{i}.png"
        if img_path.exists():
            clean_img = clean_and_crop_image(img_path, canvas_rgb, make_transparent=True)
            clean_img_path = images_dir / f"content_img_{i}_clean.png"
            clean_img.save(str(clean_img_path))
            slide.shapes.add_picture(str(clean_img_path), Inches(8.5), Inches(1.5), width=Inches(6.5))
            
        # Draw image for video
        img = Image.new('RGB', (1920, 1080), canvas_rgb)
        draw = ImageDraw.Draw(img)
        draw_wrapped_text(draw, c_data["title"], subtitle_font, 100, 60, 40, blue_rgb)
        
        y = 250
        for bullet in bullets:
            if not str(bullet).strip(): continue
            y = draw_wrapped_text(draw, str(bullet).strip().lstrip('-').strip(), body_font, 150, y, 35, text_rgb, bullet=True)
            
        if img_path.exists():
            clean_img = Image.open(str(clean_img_path)).convert("RGBA")
            clean_img.thumbnail((900, 900))
            img.paste(clean_img, (960 + 50, 1080 // 2 - clean_img.height // 2 + 50), clean_img)
            
        img.save(slide_images_dir / f"slide_{i+1}.png")
        scenes.append(Scene(c_data["spoken_text"], c_data["title"], c_data["text_left"], False))
            
    # Ending Slide
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*canvas_rgb)
    
    txBox = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(14), Inches(3))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    
    brand_site = os.getenv("BRAND_SITE", "inmibot.es")
    end_text = slides_data["ending_slide"]["spoken_text"]
    
    parts = end_text.split(brand_site)
    for idx, part in enumerate(parts):
        run = p.add_run()
        run.text = part
        run.font.size = Pt(54)
        run.font.color.rgb = RGBColor(*navy_rgb)
        run.font.bold = True
        
        if idx < len(parts) - 1:
            run_site = p.add_run()
            run_site.text = brand_site
            run_site.font.size = Pt(54)
            run_site.font.color.rgb = RGBColor(*blue_rgb)
            run_site.font.bold = True
    
    img = Image.new('RGB', (1920, 1080), canvas_rgb)
    draw = ImageDraw.Draw(img)
    
    # Custom rendering for the final slide to support multi-color text in Pillow
    lines = []
    import textwrap
    wrapped = textwrap.wrap(end_text, width=30)
    current_y = 400
    for line in wrapped:
        line_parts = line.split(brand_site)
        
        # Calculate total line width to center it
        total_w = 0
        for idx, part in enumerate(line_parts):
            bbox = draw.textbbox((0,0), part, font=subtitle_font)
            total_w += (bbox[2] - bbox[0])
            if idx < len(line_parts) - 1:
                bbox_site = draw.textbbox((0,0), brand_site, font=subtitle_font)
                total_w += (bbox_site[2] - bbox_site[0])
                
        current_x = 960 - (total_w // 2)
        line_h = 0
        
        for idx, part in enumerate(line_parts):
            bbox = draw.textbbox((0,0), part, font=subtitle_font)
            part_w = bbox[2] - bbox[0]
            line_h = max(line_h, bbox[3] - bbox[1])
            draw.text((current_x, current_y), part, font=subtitle_font, fill=navy_rgb)
            current_x += part_w
            
            if idx < len(line_parts) - 1:
                bbox_site = draw.textbbox((0,0), brand_site, font=subtitle_font)
                site_w = bbox_site[2] - bbox_site[0]
                line_h = max(line_h, bbox_site[3] - bbox_site[1])
                draw.text((current_x, current_y), brand_site, font=subtitle_font, fill=blue_rgb)
                current_x += site_w
                
        current_y += line_h + 15
        
    img.save(slide_images_dir / f"slide_{len(slides_data['content_slides']) + 1}.png")
    scenes.append(Scene(slides_data["ending_slide"]["spoken_text"], "", "", True))
    
    prs.save(pptx_path)
    return scenes

def process_article(url, lang, http):
    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(http.request("GET", url).text, "html.parser")
        root = soup.find("article") or soup.find("main")
        if not root:
            return None
        for unwanted in root.select("nav, footer, aside, script, style, form, .share, .related"):
            unwanted.decompose()
        heading = root.find("h1") or soup.find("h1")
        title = heading.get_text(" ", strip=True) if heading else ""
        parts = []
        for node in root.find_all(["h2", "h3", "p", "li"]):
            text = re.sub(r"\s+", " ", node.get_text(" ", strip=True))
            if len(text) >= 20:
                parts.append(text)
        body = "\n\n".join(parts)
        if not title or len(body) < 200:
            return None
        return f"# {title}\n\n{body}"
    except Exception as e:
        print(f"Error parsing {url}: {e}")
        return None

def process_single_url(lang, url, output_base, client, http):
    slug = urlparse(url).path.strip('/').split('/')[-1]
    if not slug:
        slug = "index"
    
    out_dir = output_base / lang / slug
    out_dir.mkdir(parents=True, exist_ok=True)
    
    pptx_path = out_dir / f"{slug}.pptx"
    video_path = out_dir / "youtube.mp4"
    meta_path = out_dir / "metadata.md"
    
    if pptx_path.exists() and video_path.exists() and meta_path.exists():
        print(f"Skipping {url}, already exists.")
        return
        
    print(f"Processing {url}")
    article_text = process_article(url, lang, http)
    if not article_text:
        print(f"Could not extract text for {url}")
        return
        
    cta = cta_es if lang == "es" else cta_en
    prompt = f"""
Generate a PowerPoint presentation structure based on this article.
The language of the presentation must be exactly the same as the article ({lang}).

Article:
{article_text}

Rules:
1. 1 title slide, 2 to 5 content slides, 1 ending slide.
2. The content slide 'text_left' should be highly useful and relevant, NO generic advice like 'have your documentation ready', focus on exactly what you will find, specific requirements, and where to book. Limit to 3-4 short bullet points.
3. The ending slide must ONLY have the spoken_text: "{cta}".
4. Generate an 'image_prompt' for the title slide and each content slide. The prompt must be in English. The image should NOT contain any text or numbers. It should be on a background color exactly matching {brand_canvas}. It should visually represent the slide.
5. Provide a spoken_text for each slide for a narration.

Respond ONLY with this JSON structure:
{{
  "title_slide": {{"title": "...", "spoken_text": "...", "image_prompt": "..."}},
  "content_slides": [
    {{"title": "...", "text_left": "...", "image_prompt": "...", "spoken_text": "..."}}
  ],
  "ending_slide": {{"spoken_text": "{cta}"}}
}}
"""
    try:
        resp = client.models.generate_content(
            model='gemini-2.5-flash',
            contents=prompt,
            config=types.GenerateContentConfig(
                response_mime_type="application/json",
            )
        )
        # Find json block
        m = re.search(r'\{.*\}', resp.text, re.DOTALL)
        if not m:
            print(f"Failed to parse JSON for {url}")
            return
        slides_data = json.loads(m.group(0))
    except Exception as e:
        print(f"Error generating text for {url}: {e}")
        return
        
    images_dir = out_dir / "assets"
    images_dir.mkdir(exist_ok=True)
    
    # Generate title image
    title_img_prompt = slides_data.get("title_slide", {}).get("image_prompt", "")
    if title_img_prompt:
        title_img_prompt += f" A prompt for an AI image generator to create an illustration ONLY. ABSOLUTELY NO TEXT, NO WORDS, NO LETTERS, NO NUMBERS inside the image. It must be a flat, clean illustration relevant to the article. The background MUST be a completely solid, uniform, untextured fill of exact hex color {brand_canvas}. Use {brand_navy} and {brand_blue} for the illustration accents. IMPORTANT: If any national flags (such as the Spanish flag) are depicted, their original real-world colors must be strictly preserved and not changed to the accent colors."
        try:
            title_img_resp = client.models.generate_content(
                model='gemini-2.5-flash-image',
                contents=title_img_prompt,
                config=types.GenerateContentConfig(
                    response_modalities=["IMAGE"],
                    image_config=types.ImageConfig(
                        aspect_ratio="3:4"
                    )
                )
            )
            title_img_path = images_dir / "title_img.png"
            title_image_bytes = None
            if title_img_resp.candidates and title_img_resp.candidates[0].content.parts:
                for part in title_img_resp.candidates[0].content.parts:
                    if part.inline_data:
                        title_image_bytes = part.inline_data.data
                        break
            if title_image_bytes:
                with open(title_img_path, "wb") as f:
                    f.write(title_image_bytes)
        except Exception as e:
            print(f"Error generating title image for {url}: {e}")
    
    for i, c_data in enumerate(slides_data.get("content_slides", [])):
        img_prompt = c_data.get("image_prompt", "") + f" A prompt for an AI image generator to create an illustration ONLY. ABSOLUTELY NO TEXT, NO WORDS, NO LETTERS, NO NUMBERS inside the image. It must be a flat, clean illustration relevant to the article. The background MUST be a completely solid, uniform, untextured fill of exact hex color {brand_canvas}. Use {brand_navy} and {brand_blue} for the illustration accents. IMPORTANT: If any national flags (such as the Spanish flag) are depicted, their original real-world colors must be strictly preserved and not changed to the accent colors."
        try:
            img_resp = client.models.generate_content(
                model='gemini-2.5-flash-image',
                contents=img_prompt,
                config=types.GenerateContentConfig(
                    response_modalities=["IMAGE"],
                    image_config=types.ImageConfig(
                        aspect_ratio="3:4"
                    )
                )
            )
            img_path = images_dir / f"content_img_{i}.png"
            image_bytes = None
            if img_resp.candidates and img_resp.candidates[0].content.parts:
                for part in img_resp.candidates[0].content.parts:
                    if part.inline_data:
                        image_bytes = part.inline_data.data
                        break
            if image_bytes:
                with open(img_path, "wb") as f:
                    f.write(image_bytes)
            else:
                print(f"No image bytes returned for slide {i} of {url}")
        except Exception as e:
            print(f"Error generating image for slide {i} of {url}: {e}")
            
    slide_images_dir = out_dir / "slides"
    slide_images_dir.mkdir(exist_ok=True)
    
    try:
        scenes = create_pptx_and_images(slides_data, images_dir, pptx_path, slide_images_dir)
    except Exception as e:
        print(f"Error creating PPTX for {url}: {e}")
        return
        
    audio_path = out_dir / "narration.mp3"
    try:
        asyncio.run(EdgeTTSProvider({"language": lang}).generate(" ".join(s.spoken_text for s in scenes), str(audio_path)))
        
        vc = VideoComposer({
            "video": {"resolution": (1920, 1080), "fps": 24}, 
            "branding": {"logo_path": os.getenv("BRAND_LOGO_PATH") or "scripts/assets/brand/logo.png"}, 
            "music": {"folder": "scripts/assets/music", "volume": 0.2, "fade_duration": 2}
        })
        vc.compose(str(audio_path), str(slide_images_dir), scenes, str(video_path))
    except Exception as e:
        print(f"Error creating video for {url}: {e}")
        
    # Metadata
    meta_prompt = f"Generate a YouTube description and LinkedIn post for the following presentation. Language: {lang}.\n\nSlides: {json.dumps(slides_data)}\n\nInclude appropriate hashtags."
    try:
        meta_resp = client.models.generate_content(
            model='gemini-2.5-flash',
            contents=meta_prompt
        )
        meta_path.write_text(meta_resp.text, encoding="utf-8")
    except Exception as e:
        print(f"Error generating metadata for {url}: {e}")
        
    print(f"Successfully processed {url}")


def main():
    try:
        client = genai.Client(vertexai=True, location="us-central1")
    except Exception as e:
        print("Failed to initialize Google Cloud. Did you run 'gcloud auth application-default login'?")
        raise e
    
    http = Http()
    sitemap_url = "https://inmibot.es/sitemap.xml"
    urls_es = discover_urls(sitemap_url, "es", http)
    urls_en = discover_urls(sitemap_url, "en", http)
    
    output_base = Path("generated/powerpoints")
    
    all_urls = [("es", u) for u in urls_es] + [("en", u) for u in urls_en]
    
    max_workers = 5
    with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = [
            executor.submit(process_single_url, lang, url, output_base, client, http) 
            for lang, url in all_urls
        ]
        concurrent.futures.wait(futures)

if __name__ == "__main__":
    main()

