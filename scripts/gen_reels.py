#!/usr/bin/env python3
import os
import sys
import json
from pathlib import Path

# Ensure we're in the right working directory or adjust sys.path
sys.path.insert(0, str(Path(__file__).parent))

# Load .env manually if python-dotenv is not installed
env_path = Path(__file__).parent.parent / ".env"
if env_path.exists():
    with open(env_path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#"):
                if "=" in line:
                    key, val = line.split("=", 1)
                    key = key.strip()
                    val = val.strip().strip("'").strip('"')
                    if key not in os.environ:
                        os.environ[key] = val

from google import genai
from google.genai import types

import numpy as np
import pptx
from PIL import Image, ImageDraw, ImageFont

# Disable ANTIALIAS error for moviepy
import PIL.Image
if not hasattr(PIL.Image, 'ANTIALIAS'):
    PIL.Image.ANTIALIAS = PIL.Image.LANCZOS

from moviepy.editor import VideoFileClip, ImageClip, CompositeVideoClip

def hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip('#')
    if len(hex_color) == 6:
        return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
    return (255, 255, 255) # fallback

def load_font(size, is_bold=False):
    font_path = os.getenv("BRAND_FONT_PATH", "")
    if font_path and os.path.exists(font_path):
        try:
            return ImageFont.truetype(font_path, size)
        except:
            pass
            
    # macOS system fonts
    mac_fonts = [
        "/System/Library/Fonts/HelveticaNeue.ttc",
        "/System/Library/Fonts/Helvetica.ttc",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf"
    ]
    
    for p in mac_fonts:
        if os.path.exists(p):
            try:
                if p.endswith('.ttc') and is_bold:
                    # For HelveticaNeue.ttc and Helvetica.ttc, index 1 is usually Bold
                    return ImageFont.truetype(p, size, index=1)
                else:
                    return ImageFont.truetype(p, size)
            except:
                continue
    
    return ImageFont.load_default()

def draw_wrapped_text(draw, text, font, text_color, max_width, x_center, start_y, align="center"):
    words = text.split()
    lines = []
    current_line = []
    
    for word in words:
        test_line = ' '.join(current_line + [word])
        bbox = draw.textbbox((0,0), test_line, font=font)
        if (bbox[2] - bbox[0]) <= max_width:
            current_line.append(word)
        else:
            if current_line:
                lines.append(' '.join(current_line))
                current_line = [word]
            else:
                lines.append(word)
                current_line = []
                
    if current_line:
        lines.append(' '.join(current_line))
        
    current_y = start_y
    for line in lines:
        bbox = draw.textbbox((0,0), line, font=font)
        w = bbox[2] - bbox[0]
        h = bbox[3] - bbox[1]
        
        if align == "center":
            x = x_center - w // 2
        elif align == "left":
            x = x_center # treated as left margin
            
        draw.text((x, current_y), line, font=font, fill=text_color)
        current_y += h + max(10, int(h * 0.3))
        
    return current_y

def extract_texts_from_pptx(pptx_path):
    prs = pptx.Presentation(pptx_path)
    slides_info = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, 'text') and shape.text.strip():
                texts.append(shape.text.strip())
        slides_info.append(texts)
    return slides_info

def generate_vertical_slide(texts, image_path, output_path, is_title=False):
    W, H = 1080, 1920
    margin = 120
    max_w = W - 2 * margin
    
    bg_color = hex_to_rgb(os.getenv("BRAND_COLOR_CANVAS", "#f8fafc"))
    navy_color = hex_to_rgb(os.getenv("BRAND_COLOR_NAVY", "#0c3d6d"))
    text_color = hex_to_rgb(os.getenv("BRAND_COLOR_TEXT", "#0f172a"))
    
    img = Image.new('RGB', (W, H), color=bg_color)
    draw = ImageDraw.Draw(img)
    
    title_font = load_font(85 if is_title else 70, is_bold=True)
    body_font = load_font(48, is_bold=False)
    
    if is_title:
        title_text = texts[0] if texts else "Presentation"
        y_text = 300
        
        y_end = draw_wrapped_text(draw, title_text, title_font, navy_color, max_w, W//2, y_text, align="center")
        
        if image_path and os.path.exists(image_path):
            try:
                slide_img = Image.open(image_path).convert("RGBA")
                target_w = 850
                target_h = int(slide_img.height * (target_w / slide_img.width))
                slide_img = slide_img.resize((target_w, target_h), Image.LANCZOS)
                
                y_img = max(y_end + 150, (H + y_end - target_h) // 2)
                img.paste(slide_img, ((W-target_w)//2, y_img), slide_img)
            except Exception as e:
                print(f"Error loading image {image_path}: {e}")
                
    elif not image_path and len(texts) == 1:
        # Ending slide
        text = texts[0]
        y_text = H // 2 - 150
        draw_wrapped_text(draw, text, title_font, navy_color, max_w, W//2, y_text, align="center")
        
    else:
        # Content Slide
        title_text = texts[0] if texts else ""
        body_text = texts[1] if len(texts) > 1 else ""
        
        y_text = 150
        y_end = draw_wrapped_text(draw, title_text, title_font, navy_color, max_w, W//2, y_text, align="center")
        
        if image_path and os.path.exists(image_path):
            try:
                slide_img = Image.open(image_path).convert("RGBA")
                target_h = 550
                target_w = int(slide_img.width * (target_h / slide_img.height))
                if target_w > max_w:
                    target_w = max_w
                    target_h = int(slide_img.height * (target_w / slide_img.width))
                    
                slide_img = slide_img.resize((target_w, target_h), Image.LANCZOS)
                
                y_img = y_end + 80
                img.paste(slide_img, ((W-target_w)//2, y_img), slide_img)
                y_end = y_img + target_h + 80
            except Exception as e:
                print(f"Error loading image {image_path}: {e}")
        else:
            y_end += 100
            
        y_body = y_end
        
        for bullet in body_text.split('\n'):
            bullet = bullet.strip()
            if not bullet: continue
            
            # bullet indicator
            draw.text((margin, y_body), "•", font=body_font, fill=text_color)
            
            # bullet text
            bullet_margin = margin + 50
            bullet_max_w = W - margin - bullet_margin
            y_body = draw_wrapped_text(draw, bullet, body_font, text_color, bullet_max_w, bullet_margin, y_body, align="left")
            y_body += 30

    img.save(output_path)

def get_slide_timings(video_path, num_slides):
    clip = VideoFileClip(str(video_path))
    dur = clip.duration
    
    times = np.arange(0, dur, 0.5)
    frames = [clip.get_frame(t) for t in times]
    
    slide_starts = [0.0]
    for i in range(1, len(frames)):
        diff = np.mean(np.abs(frames[i].astype(float) - frames[i-1].astype(float)))
        if diff > 10.0:
            t = times[i]
            if t - slide_starts[-1] > 2.0:
                slide_starts.append(t)
    
    if len(slide_starts) > num_slides:
        slide_starts = slide_starts[:num_slides]
    elif len(slide_starts) < num_slides:
        missing = num_slides - len(slide_starts)
        time_left = dur - slide_starts[-1]
        avg_time = time_left / (missing + 1)
        for _ in range(missing):
            slide_starts.append(slide_starts[-1] + avg_time)
            
    clip.close()
    return slide_starts, dur

import imageio_ffmpeg
import subprocess

def process_reel(folder_path, client, only_metadata=False):
    folder = Path(folder_path)
    youtube_mp4 = folder / "youtube.mp4"
    if not youtube_mp4.exists():
        return

    lang = folder.parent.name
    article = folder.name
    target_folder = Path("generated/reels") / lang / article
    target_folder.mkdir(parents=True, exist_ok=True)

    reel_mp4 = target_folder / "reel.mp4"
    reel_metadata_json = target_folder / "reels_metadata.json"

    if not only_metadata and not reel_mp4.exists():
        print(f"Generating Reel video for {folder.name}")
        
        pptx_files = list(folder.glob("*.pptx"))
        if not pptx_files:
            print(f"No pptx found in {folder}")
            return
        slides_info = extract_texts_from_pptx(pptx_files[0])
        
        starts, dur = get_slide_timings(youtube_mp4, len(slides_info))
        
        assets_dir = folder / "assets"
        temp_slides_dir = target_folder / "temp_slides"
        temp_slides_dir.mkdir(exist_ok=True)
        
        visual_clips = []
        for i, texts in enumerate(slides_info):
            img_out = temp_slides_dir / f"slide_{i}.png"
            
            if i == 0:
                img_src_clean = assets_dir / "title_img_clean.png"
                img_src = img_src_clean if img_src_clean.exists() else assets_dir / "title_img.png"
                generate_vertical_slide(texts, str(img_src) if img_src.exists() else None, str(img_out), is_title=True)
            elif i == len(slides_info) - 1:
                generate_vertical_slide(texts, None, str(img_out), is_title=False)
            else:
                img_src_clean = assets_dir / f"content_img_{i-1}_clean.png"
                img_src = img_src_clean if img_src_clean.exists() else assets_dir / f"content_img_{i-1}.png"
                generate_vertical_slide(texts, str(img_src) if img_src.exists() else None, str(img_out), is_title=False)
                
            start_time = starts[i]
            end_time = starts[i+1] if i + 1 < len(starts) else dur
            
            # To fix crossfade black flashes, we extend the duration of this clip by 0.5s if it's not the last one
            clip_dur = (end_time - start_time) + (0.5 if i < len(slides_info) - 1 else 0)
            
            slide_clip = ImageClip(str(img_out)).set_start(start_time).set_duration(clip_dur)
            if i > 0:
                slide_clip = slide_clip.crossfadein(0.5)
            visual_clips.append(slide_clip)
            
        final_video = CompositeVideoClip(visual_clips, size=(1080, 1920)).set_duration(dur)
        
        temp_video_mp4 = target_folder / "temp_video.mp4"
        
        final_video.write_videofile(
            str(temp_video_mp4),
            codec="libx264",
            fps=24,
            preset="ultrafast",
            audio=False,
            logger=None
        )
        
        final_video.close()
        
        # Mux the video and cleanly rebuild the audio mix with ffmpeg to avoid moviepy AAC corruption
        ffmpeg_exe = imageio_ffmpeg.get_ffmpeg_exe()
        narration_mp3 = folder / "narration.mp3"
        music_dir = Path(__file__).parent / "assets" / "music"
        music_files = list(music_dir.glob("*.mp3"))
        music_mp3 = music_files[0] if music_files else None

        if narration_mp3.exists() and music_mp3:
            mux_cmd = [
                ffmpeg_exe,
                "-y",
                "-i", str(temp_video_mp4),
                "-i", str(narration_mp3),
                "-stream_loop", "-1",
                "-i", str(music_mp3),
                "-filter_complex",
                "[2:a]volume=0.2[bgm];[1:a][bgm]amix=inputs=2:duration=first:dropout_transition=2[a]",
                "-map", "0:v:0",
                "-map", "[a]",
                "-c:v", "copy",
                "-c:a", "aac",
                "-b:a", "192k",
                "-ac", "2",
                "-shortest",
                str(reel_mp4)
            ]
        else:
            # Fallback to direct stream copy from youtube.mp4 (though it might be corrupt at the end)
            mux_cmd = [
                ffmpeg_exe,
                "-y",
                "-i", str(temp_video_mp4),
                "-i", str(youtube_mp4),
                "-c:v", "copy",
                "-c:a", "copy",
                "-map", "0:v:0",
                "-map", "1:a:0",
                "-shortest",
                str(reel_mp4)
            ]
        
        try:
            subprocess.run(mux_cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        except Exception as e:
            print(f"Error muxing audio for {folder.name}: {e}")
            
        # --- NEW CODE FOR YOUTUBE SPECIFIC REEL ---
        reel_youtube_mp4 = target_folder / "reel_youtube.mp4"
        yt_music_dir = Path(__file__).parent / "assets" / "youtube_music"
        yt_music_files = list(yt_music_dir.glob("*.mp3")) if yt_music_dir.exists() else []
        yt_music_mp3 = yt_music_files[0] if yt_music_files else None

        if narration_mp3.exists():
            if yt_music_mp3:
                mux_yt_cmd = [
                    ffmpeg_exe,
                    "-y",
                    "-i", str(temp_video_mp4),
                    "-i", str(narration_mp3),
                    "-stream_loop", "-1",
                    "-i", str(yt_music_mp3),
                    "-filter_complex",
                    "[2:a]volume=0.2[bgm];[1:a][bgm]amix=inputs=2:duration=first:dropout_transition=2[a]",
                    "-map", "0:v:0",
                    "-map", "[a]",
                    "-c:v", "copy",
                    "-c:a", "aac",
                    "-b:a", "192k",
                    "-ac", "2",
                    "-shortest",
                    str(reel_youtube_mp4)
                ]
            else:
                # Just voiceover
                mux_yt_cmd = [
                    ffmpeg_exe,
                    "-y",
                    "-i", str(temp_video_mp4),
                    "-i", str(narration_mp3),
                    "-c:v", "copy",
                    "-c:a", "aac",
                    "-b:a", "192k",
                    "-map", "0:v:0",
                    "-map", "1:a:0",
                    "-shortest",
                    str(reel_youtube_mp4)
                ]
            try:
                subprocess.run(mux_yt_cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            except Exception as e:
                print(f"Error muxing youtube audio for {folder.name}: {e}")
        # ------------------------------------------

        if temp_video_mp4.exists():
            temp_video_mp4.unlink()

            
        for f in temp_slides_dir.iterdir():
            f.unlink()
        temp_slides_dir.rmdir()
        print(f"Successfully generated {reel_mp4}")

    if only_metadata or not reel_metadata_json.exists():
        print(f"Generating Social Media Metadata for {folder.name}")
        md_content = ""
        metadata_md = folder / "metadata.md"
        metadata_json = folder / "metadata.json"
        
        if metadata_md.exists():
            md_content = metadata_md.read_text(encoding="utf-8")
        elif metadata_json.exists():
            md_content = metadata_json.read_text(encoding="utf-8")
        else:
            print(f"No metadata source found for {folder.name}")
            return
            
        prompt = f"""
You are an expert social media manager.
Read the following content (which may be a YouTube description or existing metadata for a video).
Generate engaging descriptions and tags optimized for TikTok, Facebook, Instagram Reels, YouTube Shorts, and LinkedIn.
Return the result as a strict JSON with the exact following schema:

{{
    "tiktok": {{
        "description": "Engaging description for TikTok (excluding tags)",
        "tags": ["tag1", "tag2"]
    }},
    "facebook": {{
        "description": "Engaging description for Facebook (excluding tags)",
        "tags": ["tag1", "tag2"]
    }},
    "instagram": {{
        "description": "Engaging description for Instagram Reels (excluding tags)",
        "tags": ["tag1", "tag2"]
    }},
    "youtube_shorts": {{
        "title": "Engaging title for YouTube Shorts (under 100 characters)",
        "description": "Engaging description for YouTube Shorts (excluding tags)",
        "tags": ["tag1", "tag2"]
    }},
    "linkedin": {{
        "post": "Professional and engaging post for LinkedIn (excluding tags)",
        "tags": ["tag1", "tag2"]
    }}
}}

Make sure to extract hashtags into the 'tags' arrays without the '#' symbol.

Content:
{md_content}
"""
        try:
            resp = client.models.generate_content(
                model='gemini-2.5-flash',
                contents=prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    temperature=0.7
                ),
            )
            data = json.loads(resp.text)
            reel_metadata_json.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
            print(f"Successfully generated metadata: {reel_metadata_json}")
        except Exception as e:
            print(f"Failed to generate metadata for {folder.name}: {e}")

def main():
    import argparse
    parser = argparse.ArgumentParser(description="Generate reels and metadata")
    parser.add_argument("--only-metadata", action="store_true", help="Only re-make the metadata, skip video generation")
    args = parser.parse_args()

    client = genai.Client(vertexai=True, location="us-central1")
    base_dir = Path("generated/powerpoints")
    
    import concurrent.futures
    
    folders_to_process = []
    for lang_dir in base_dir.iterdir():
        if lang_dir.is_dir():
            for article_dir in lang_dir.iterdir():
                if article_dir.is_dir() and (article_dir / "youtube.mp4").exists():
                    folders_to_process.append(article_dir)
                    
    print(f"Found {len(folders_to_process)} folders to process.")
    
    # Use more workers if we're only doing metadata (API bound) vs video processing (CPU bound)
    workers = 10 if args.only_metadata else 7
    
    with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as executor:
        futures = [executor.submit(process_reel, folder, client, args.only_metadata) for folder in folders_to_process]
        for future in concurrent.futures.as_completed(futures):
            try:
                future.result()
            except Exception as e:
                print(f"Error in task: {e}")

if __name__ == "__main__":
    main()
